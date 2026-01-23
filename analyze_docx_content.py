import docx
import sys
import os

def read_docx(path):
    if not os.path.exists(path):
        return f"File not found: {path}"
    try:
        doc = docx.Document(path)
        full_text = []
        full_text.append(f"=== READING: {os.path.basename(path)} ===")
        
        # Read Paragraphs
        for p in doc.paragraphs:
            if p.text.strip():
                full_text.append(p.text)
        
        # Read Tables (CRITICAL for this analysis)
        full_text.append("\n--- TABLES ---\n")
        for i, table in enumerate(doc.tables):
            full_text.append(f"Table {i+1}:")
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells]
                full_text.append(" | ".join(row_text))
            full_text.append("") 
            
        return "\n".join(full_text)
    except Exception as e:
        return f"Error reading {path}: {e}"

if __name__ == "__main__":
    files = [
        "rfq_downloads/2026-01-20/c54895f0cf4c40e59671a2f4c7ec117c_RFQ_SERVICE.docx",
        "rfq_downloads/2026-01-20/0ca60069c6b440f9ad573995205787aa_RFQ_PRODUCT.docx",
        "rfq_downloads/2026-01-20/510d2187c9ea40a9acbdc7462221ed98_RFQ_SERVICE.docx",
        "rfq_downloads/2026-01-20/ce00f362685b4e9eb2096d3765f5e7e4_RFQ_SERVICE.docx",
        "rfq_downloads/2026-01-17/f3d327c5ae8341f0a54658b1283ce4c7_RFQ_PRODUCT.docx"
    ]
    
    for f in files:
        print(read_docx(f))
        print("\n" + "="*50 + "\n")
